"""Excel parser implementation (.xlsx via openpyxl)."""
from __future__ import annotations

from io import BytesIO
from typing import Union

from openpyxl import load_workbook

from peter_parser_core import BaseParser, ParsedDocument, ParserError
from peter_parser_core.common.types import ContentModel, Page, Table, TableCell


class ExcelParser(BaseParser):
    """Parse .xlsx files into ParsedDocument (sheet → Page, cells → Table)."""

    def parse(self, document: Union[bytes, str], **kwargs) -> ParsedDocument:
        """Parse Excel file into ParsedDocument.

        Args:
            document: .xlsx file as bytes or path (str).

        Returns:
            ParsedDocument with pages (one per sheet), each with a Table of cells.
        """
        try:
            if isinstance(document, str):
                wb = load_workbook(document, read_only=True, data_only=True)
            else:
                wb = load_workbook(BytesIO(document), read_only=True, data_only=True)
        except Exception as e:
            raise ParserError(f"Failed to load Excel: {e}") from e

        try:
            pages: list[Page] = []
            all_text_parts: list[str] = []
            sheet_names: list[str] = []

            for sheet_idx, sheet in enumerate(wb.worksheets):
                page_number = sheet_idx + 1
                sheet_name = sheet.title
                sheet_names.append(sheet_name)

                cells: list[TableCell] = []
                row_texts: list[str] = []

                for row_idx, row in enumerate(sheet.iter_rows()):
                    row_values: list[str] = []
                    for col_idx, cell in enumerate(row):
                        val = cell.value
                        content = "" if val is None else str(val).strip()
                        cells.append(
                            TableCell(
                                row=row_idx,
                                column=col_idx,
                                content=content,
                                row_span=1,
                                column_span=1,
                            )
                        )
                        row_values.append(content)
                    row_texts.append("\t".join(row_values))

                max_row = max((c.row for c in cells), default=0) + 1
                max_col = max((c.column for c in cells), default=0) + 1

                table = Table(
                    table_index=0,
                    page_number=page_number,
                    rows=max_row,
                    columns=max_col,
                    cells=cells,
                )
                sheet_text = "\n".join(row_texts)
                all_text_parts.append(sheet_text)

                page = Page(
                    page_number=page_number,
                    text=sheet_text,
                    tables=[table],
                    images=[],
                )
                pages.append(page)
        finally:
            wb.close()

        full_text = "\n\n".join(all_text_parts)
        content = ContentModel(text=full_text)
        metadata = {
            "source": "excel",
            "sheet_names": sheet_names,
            "num_sheets": len(pages),
        }
        return ParsedDocument(
            pages=pages,
            elements=[],
            content=content,
            metadata=metadata,
        )
