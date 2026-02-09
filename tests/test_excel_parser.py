"""Tests for ExcelParser."""
from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pytest

from openpyxl import Workbook

from peter_parser.impl.parser.excel import ExcelParser


def _create_sample_xlsx_bytes() -> bytes:
    """Create minimal .xlsx in memory."""
    wb = Workbook()
    ws = wb.active
    if ws:
        ws.title = "Sheet1"
        ws["A1"] = "Name"
        ws["B1"] = "Value"
        ws["A2"] = "foo"
        ws["B2"] = 100
        ws["A3"] = "bar"
        ws["B3"] = 200
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _create_sample_xlsx_with_two_sheets() -> bytes:
    """Create .xlsx with two sheets."""
    wb = Workbook()
    ws1 = wb.active
    if ws1:
        ws1.title = "Summary"
        ws1["A1"] = "Total"
        ws1["B1"] = 300
    ws2 = wb.create_sheet("Details")
    ws2["A1"] = "ID"
    ws2["B1"] = "Amount"
    ws2["A2"] = 1
    ws2["B2"] = 100
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_excel_parser_parse_bytes():
    """ExcelParser.parse returns ParsedDocument from bytes."""
    parser = ExcelParser()
    xlsx_bytes = _create_sample_xlsx_bytes()
    doc = parser.parse(xlsx_bytes)
    assert doc is not None
    assert len(doc.pages) == 1
    assert doc.pages[0].page_number == 1
    assert len(doc.pages[0].tables) == 1
    table = doc.pages[0].tables[0]
    assert table.rows >= 2
    assert table.columns >= 2
    assert len(table.cells) >= 4
    assert doc.content.text
    assert doc.metadata.get("source") == "excel"
    assert doc.metadata.get("num_sheets") == 1


def test_excel_parser_parse_two_sheets():
    """ExcelParser.parse handles multiple sheets."""
    parser = ExcelParser()
    xlsx_bytes = _create_sample_xlsx_with_two_sheets()
    doc = parser.parse(xlsx_bytes)
    assert len(doc.pages) == 2
    assert doc.pages[0].page_number == 1
    assert doc.pages[1].page_number == 2
    assert doc.metadata.get("num_sheets") == 2
    assert doc.metadata.get("sheet_names") == ["Summary", "Details"]


def test_excel_parser_parse_path(tmp_path: Path):
    """ExcelParser.parse accepts file path."""
    xlsx_bytes = _create_sample_xlsx_bytes()
    path = tmp_path / "sample.xlsx"
    path.write_bytes(xlsx_bytes)
    parser = ExcelParser()
    doc = parser.parse(str(path))
    assert len(doc.pages) == 1
    assert doc.metadata.get("source") == "excel"


def test_excel_parser_invalid_bytes_raises():
    """ExcelParser.parse raises ParserError for invalid input."""
    from peter_parser_core import ParserError

    parser = ExcelParser()
    with pytest.raises(ParserError):
        parser.parse(b"not excel content")
